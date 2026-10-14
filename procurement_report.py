"""Standalone procurement CSV -> Excel report. Run: python procurement_report.py"""
from pathlib import Path
from datetime import date, datetime
import argparse
try:
    import pandas as pd
except ModuleNotFoundError as exc:
    if exc.name != 'pandas':
        raise
    raise SystemExit(
        "Missing dependency 'pandas'. Install the project dependencies with: "
        "python -m pip install -r requirements.txt"
    ) from exc
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
DATE_COLUMNS = ('[REQ]Requisition Date (Date)', '[APS (REQ)] Assigned Date')

def find_column(columns, fragment):
    matches = [c for c in columns if fragment.lower() in c.lower()]
    if len(matches) != 1:
        raise ValueError(f"Expected exactly one column containing {fragment!r}, found: {matches}")
    return matches[0]

def parse_us_dates(series, name):
    raw = series.astype('string').str.strip()
    dates = pd.to_datetime(raw, format='%m-%d-%Y', errors='coerce')
    # Also allow US dates with slashes; never infer DMY.
    fallback = pd.to_datetime(raw, format='%m/%d/%Y', errors='coerce')
    dates = dates.fillna(fallback)
    invalid = raw.notna() & raw.ne('') & dates.isna()
    return dates, invalid

def run(source, destination, as_of):
    df = pd.read_csv(source, dtype=str, encoding='utf-8-sig', sep=None, engine='python', keep_default_na=False)
    df.columns = df.columns.str.strip()
    assigned_col = find_column(df.columns, 'Assigned Date')
    req_col = find_column(df.columns, 'Requisition Date')
    parsed_req, bad_req = parse_us_dates(df[req_col], req_col)
    parsed_assigned, bad_assigned = parse_us_dates(df[assigned_col], assigned_col)
    df[req_col] = parsed_req
    df[assigned_col] = parsed_assigned
    today = pd.Timestamp(as_of)
    days = (today - parsed_assigned).dt.days
    df['Days Since Assigned'] = days.astype('Int64')
    df['Greater than 2 days'] = days.map(lambda d: 'Yes' if d > 2 else 'No' if pd.notna(d) else 'Review')
    df['Data Quality'] = [
        'Invalid/missing date' if bad_req.iloc[i] or bad_assigned.iloc[i] or pd.isna(parsed_assigned.iloc[i]) or pd.isna(parsed_req.iloc[i])
        else 'Future assigned date' if parsed_assigned.iloc[i] > today
        else 'Assigned before requisition' if parsed_assigned.iloc[i] < parsed_req.iloc[i]
        else 'OK'
        for i in range(len(df))
    ]
    # Preserve original text on an audit sheet for traceability.
    original = pd.read_csv(source, dtype=str, encoding='utf-8-sig', sep=None, engine='python', keep_default_na=False)
    unique_pr = next((c for c in df if c.strip().lower() == 'ref pr'), None)
    summary = pd.DataFrame([
        ('Report as of', as_of.isoformat()),
        ('Total rows', len(df)),
        ('Unique PRs', df[unique_pr].replace('', pd.NA).nunique() if unique_pr else 'N/A'),
        ('Rows over 2 days', int((df['Greater than 2 days'] == 'Yes').sum())),
        ('Rows 2 days or less', int((df['Greater than 2 days'] == 'No').sum())),
        ('Rows needing review', int((df['Data Quality'] != 'OK').sum())),
    ], columns=['Metric', 'Value'])
    destination.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(destination, engine='openpyxl', datetime_format='DD-MM-YYYY', date_format='DD-MM-YYYY') as writer:
        summary.to_excel(writer, sheet_name='Summary', index=False)
        df.to_excel(writer, sheet_name='Approval Report', index=False)
        original.to_excel(writer, sheet_name='Raw CSV', index=False)
    wb = load_workbook(destination)
    for ws in wb.worksheets:
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 30
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='17365D')
            cell.font = Font(bold=True, color='FFFFFF')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            max_len = max((len(str(c.value)) for c in list(col)[:300] if c.value is not None), default=12)
            ws.column_dimensions[letter].width = min(42, max(15, max_len + 2))
    report = wb['Approval Report']
    for j, cell in enumerate(report[1], 1):
        if cell.value in (req_col, assigned_col):
            for row in report.iter_rows(min_row=2, min_col=j, max_col=j):
                row[0].number_format = 'DD-MM-YYYY'
    wb.save(destination)
    return len(df)

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Convert procurement CSV into an approval Excel report.')
    parser.add_argument('--input', type=Path, help='Input CSV; defaults to newest CSV in input/')
    parser.add_argument('--output', type=Path, help='Output XLSX path')
    parser.add_argument('--as-of', type=date.fromisoformat, default=date.today(), help='Report date YYYY-MM-DD; default system today')
    args = parser.parse_args()
    candidates = sorted((BASE / 'input').glob('*.csv'), key=lambda p: p.stat().st_mtime, reverse=True)
    source = args.input or (candidates[0] if candidates else None)
    if source is None:
        parser.error('No CSV found. Put a CSV into the input folder or pass --input.')
    output = args.output or BASE / 'output' / f'procurement_report_{args.as_of:%Y%m%d}.xlsx'
    count = run(source, output, args.as_of)
    print(f'Success: processed {count} rows. Report: {output}')
